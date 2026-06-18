"""File loaders that return clean text plus section metadata.

Each loader returns a ``Document`` dict::

    {
        "text": str,                 # full document text
        "file_type": str,            # pdf | docx | pptx | xlsx | csv | txt | md
        "source_file": str,
        "filename": str,
        "document_title": str | None,
        "sections": [ {text, page_number, slide_index, section_heading,
                       char_start, char_end}, ... ],
    }

One library per format, no LangChain / unstructured.
"""

import csv
import re
from pathlib import Path
from typing import Dict, List


def _section(text, *, page_number=None, slide_index=None, heading=None, start=0):
    return {
        "text": text,
        "page_number": page_number,
        "slide_index": slide_index,
        "section_heading": heading,
        "char_start": start,
        "char_end": start + len(text),
    }


def _document(text, file_type, path: Path, sections, title=None) -> Dict:
    if not sections:
        sections = [_section(text)]
    return {
        "text": text,
        "file_type": file_type,
        "source_file": str(path),
        "filename": path.name,
        "document_title": title,
        "sections": sections,
    }


def read_pdf(path: Path) -> Dict:
    try:
        import fitz  # PyMuPDF
    except Exception as exc:  # pragma: no cover - dependency guard
        raise RuntimeError("PyMuPDF is required to read PDF files (pip install pymupdf)") from exc

    doc = fitz.open(str(path))
    try:
        title = (doc.metadata or {}).get("title") or None
        sections: List[Dict] = []
        parts: List[str] = []
        offset = 0
        for page_index, page in enumerate(doc, start=1):
            text = (page.get_text() or "").strip()
            if not text:
                continue
            parts.append(text)
            sections.append(_section(text, page_number=page_index, start=offset))
            offset += len(text) + 2
    finally:
        doc.close()

    return _document("\n\n".join(parts), "pdf", path, sections, title)


def read_docx(path: Path) -> Dict:
    try:
        import docx
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("python-docx is required to read DOCX files") from exc

    document = docx.Document(str(path))
    sections: List[Dict] = []
    heading = None
    title = None
    buffer: List[str] = []
    offset = 0

    def flush():
        nonlocal offset, buffer
        if buffer:
            text = "\n\n".join(buffer)
            sections.append(_section(text, heading=heading, start=offset))
            offset += len(text) + 2
            buffer = []

    for paragraph in document.paragraphs:
        text = (paragraph.text or "").strip()
        if not text:
            continue
        style = getattr(getattr(paragraph, "style", None), "name", "") or ""
        if style.lower().startswith(("heading", "title")):
            flush()
            heading = text
            title = title or text
        else:
            buffer.append(text)
    flush()

    full_text = "\n\n".join(s["text"] for s in sections)
    return _document(full_text, "docx", path, sections, title)


def read_pptx(path: Path) -> Dict:
    try:
        from pptx import Presentation
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("python-pptx is required to read PPTX files") from exc

    presentation = Presentation(str(path))
    sections: List[Dict] = []
    parts: List[str] = []
    offset = 0
    for slide_index, slide in enumerate(presentation.slides, start=1):
        texts = [
            shape.text.strip()
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text and shape.text.strip()
        ]
        slide_text = "\n\n".join(texts).strip()
        if not slide_text:
            continue
        parts.append(slide_text)
        sections.append(_section(slide_text, slide_index=slide_index, start=offset))
        offset += len(slide_text) + 2

    return _document("\n\n".join(parts), "pptx", path, sections)


def read_xlsx(path: Path) -> Dict:
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required to read XLSX files") from exc

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    sections: List[Dict] = []
    parts: List[str] = []
    offset = 0
    for worksheet in workbook.worksheets:
        rows = []
        for row in worksheet.iter_rows(values_only=True):
            # Drop empty/blank cells so only meaningful values are ingested.
            cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
            if cells:
                rows.append("\t".join(cells))
        sheet_text = "\n".join(rows).strip()
        if not sheet_text:
            continue
        parts.append(sheet_text)
        sections.append(_section(sheet_text, heading=worksheet.title, start=offset))
        offset += len(sheet_text) + 2
    workbook.close()

    return _document("\n\n".join(parts), "xlsx", path, sections)


def read_csv(path: Path) -> Dict:
    rows = None
    for encoding in ("utf-8", "utf-8-sig", "latin1", "cp1252"):
        try:
            with path.open("r", newline="", encoding=encoding) as fh:
                # Drop empty cells and blank rows so only meaningful values are ingested.
                parsed = [
                    "\t".join(cell.strip() for cell in row if cell and cell.strip())
                    for row in csv.reader(fh)
                ]
            rows = [row for row in parsed if row]
            break
        except UnicodeDecodeError:
            continue
    if rows is None:
        raise RuntimeError(f"Failed to read CSV {path} with common encodings")

    text = "\n".join(rows)
    return _document(text, "csv", path, [_section(text)])


def read_text_or_md(path: Path) -> Dict:
    text = path.read_text(encoding="utf-8")
    file_type = "md" if path.suffix.lower() == ".md" else "txt"

    # Split on blank lines into paragraph sections.
    sections: List[Dict] = []
    offset = 0
    for block in re.split(r"\n\s*\n", text):
        block = block.strip()
        if not block:
            continue
        sections.append(_section(block, start=offset))
        offset += len(block) + 2

    full_text = "\n\n".join(s["text"] for s in sections) if sections else text
    return _document(full_text, file_type, path, sections)


_LOADERS = {
    ".pdf": read_pdf,
    ".docx": read_docx,
    ".doc": read_docx,
    ".pptx": read_pptx,
    ".ppt": read_pptx,
    ".xlsx": read_xlsx,
    ".csv": read_csv,
    ".txt": read_text_or_md,
    ".md": read_text_or_md,
}


def load_file(path: str) -> Dict:
    """Load a supported file into a ``Document`` dict (see module docstring)."""
    p = Path(path)
    loader = _LOADERS.get(p.suffix.lower())
    if loader is None:
        raise RuntimeError(f"Unsupported file type: {p.suffix}")
    return loader(p)
