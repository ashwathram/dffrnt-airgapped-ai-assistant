"""File loaders that return clean text plus rich metadata.

Each loader returns a dict with full text, file type, source path, and
section metadata for downstream chunking.
"""

from pathlib import Path
from typing import Dict, List
import re


def _normalize_whitespace(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip()


def read_pdf(path: Path) -> Dict:
    try:
        # Prefer the maintained `pypdf` package if available (mitigates known PyPDF2 CVEs)
        import pypdf as _pypdf  # type: ignore

        PdfReaderClass = _pypdf.PdfReader
    except Exception:
        try:
            import PyPDF2 as _pyPdf2

            PdfReaderClass = _pyPdf2.PdfReader
        except Exception as e:
            raise RuntimeError(
                (
                    "A PDF reader is required to read PDF files. Install with "
                    "pip install pypdf or pip install pypdf2"
                )
            ) from e

    reader = PdfReaderClass(str(path))
    sections: List[Dict] = []
    full_parts: List[str] = []
    offset = 0
    title = None
    # try to get document title from metadata
    try:
        meta = reader.metadata
        if meta and getattr(meta, "title", None):
            title = _normalize_whitespace(meta.title or "")
    except Exception:
        title = None

    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        text = text.strip()
        if not text:
            continue
        full_parts.append(text)
        start = offset
        end = offset + len(text)
        sections.append(
            {
                "text": text,
                "page_number": i,
                "slide_index": None,
                "section_heading": None,
                "char_start": start,
                "char_end": end,
            }
        )
        offset = end + 2

    full_text = "\n\n".join(full_parts)
    return {
        "text": full_text,
        "file_type": "pdf",
        "source_file": str(path),
        "filename": path.name,
        "document_title": title,
        "sections": sections
        or [
            {
                "text": full_text,
                "page_number": None,
                "slide_index": None,
                "section_heading": None,
                "char_start": 0,
                "char_end": len(full_text),
            }
        ],
    }


def read_docx(path: Path) -> Dict:
    try:
        import docx
    except Exception as e:
        raise RuntimeError(
            "python-docx is required to read DOCX files. Install with `pip install python-docx`"
        ) from e

    doc = docx.Document(str(path))
    paragraphs = [p for p in doc.paragraphs if p.text is not None]
    sections: List[Dict] = []
    current_heading = None
    current_parts: List[str] = []
    offset = 0
    title = None

    for p in paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = getattr(p, "style", None)
        style_name = getattr(style, "name", "") if style is not None else ""
        is_heading = style_name.lower().startswith("heading") or style_name.lower().startswith(
            "title"
        )
        if is_heading:
            # flush current
            if current_parts:
                full = "\n\n".join(current_parts)
                sections.append(
                    {
                        "text": full,
                        "page_number": None,
                        "slide_index": None,
                        "section_heading": current_heading,
                        "char_start": offset,
                        "char_end": offset + len(full),
                    }
                )
                offset += len(full) + 2
                current_parts = []
            current_heading = text
            if title is None:
                title = text
            continue

        current_parts.append(text)

    if current_parts:
        full = "\n\n".join(current_parts)
        sections.append(
            {
                "text": full,
                "page_number": None,
                "slide_index": None,
                "section_heading": current_heading,
                "char_start": offset,
                "char_end": offset + len(full),
            }
        )

    full_text = "\n\n".join([s["text"] for s in sections])
    if not sections:
        full_text = "\n\n".join([p.text for p in paragraphs])
        sections = [
            {
                "text": full_text,
                "page_number": None,
                "slide_index": None,
                "section_heading": None,
                "char_start": 0,
                "char_end": len(full_text),
            }
        ]

    return {
        "text": full_text,
        "file_type": "docx",
        "source_file": str(path),
        "filename": path.name,
        "document_title": title,
        "sections": sections,
    }


def read_pptx(path: Path) -> Dict:
    try:
        from pptx import Presentation
    except Exception as e:
        raise RuntimeError(
            "python-pptx is required to read PPTX files. Install with `pip install python-pptx`"
        ) from e

    prs = Presentation(str(path))
    sections: List[Dict] = []
    full_parts: List[str] = []
    offset = 0
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                t = shape.text
                if t:
                    texts.append(t.strip())
        slide_text = "\n\n".join(texts).strip()
        if not slide_text:
            continue
        full_parts.append(slide_text)
        sections.append(
            {
                "text": slide_text,
                "page_number": None,
                "slide_index": i,
                "section_heading": None,
                "char_start": offset,
                "char_end": offset + len(slide_text),
            }
        )
        offset += len(slide_text) + 2

    full_text = "\n\n".join(full_parts)
    return {
        "text": full_text,
        "file_type": "pptx",
        "source_file": str(path),
        "filename": path.name,
        "document_title": None,
        "sections": sections
        or [
            {
                "text": full_text,
                "page_number": None,
                "slide_index": None,
                "section_heading": None,
                "char_start": 0,
                "char_end": len(full_text),
            }
        ],
    }


def read_csv(path: Path) -> Dict:
    try:
        import pandas as pd
    except Exception as e:
        raise RuntimeError(
            "pandas is required to read CSV files. Install with `pip install pandas`"
        ) from e

    df = None
    last_error = None
    for read_kwargs in (
        {"encoding": "utf-8"},
        {"encoding": "utf-8-sig"},
        {"encoding": "latin1"},
        {"encoding": "cp1252"},
    ):
        try:
            df = pd.read_csv(path, **read_kwargs)
            break
        except Exception as exc:
            last_error = exc

    if df is None:
        raise RuntimeError(f"Failed to read CSV file {path} with common encodings") from last_error

    # Convert dataframe to text: header + rows
    parts = ["\t".join(map(str, df.columns))]
    for _, row in df.iterrows():
        parts.append("\t".join([str(x) for x in row.tolist()]))
    full_text = "\n".join(parts)
    sections = [
        {
            "text": full_text,
            "page_number": None,
            "slide_index": None,
            "section_heading": None,
            "char_start": 0,
            "char_end": len(full_text),
        }
    ]
    return {
        "text": full_text,
        "file_type": "csv",
        "source_file": str(path),
        "filename": path.name,
        "document_title": None,
        "sections": sections,
    }


def read_text_or_md(path: Path) -> Dict:
    text = path.read_text(encoding="utf-8")
    # split markdown headings as sections
    sections: List[Dict] = []
    # fallback: split on blank lines
    if path.suffix.lower() == ".md":
        # simple split by headings
        # naive: if headings present, split into sections by heading lines
        headings = re.findall(r"(?m:^(#{1,6})\s+(.*)$)", text)
        if headings:
            # split by lines starting with #
            blocks = re.split(r"(?m:^#{1,6}\s+.*$)", text)
            cursor = 0
            for b in blocks:
                b = b.strip()
                if not b:
                    continue
                start = cursor
                end = start + len(b)
                sections.append(
                    {
                        "text": b,
                        "page_number": None,
                        "slide_index": None,
                        "section_heading": None,
                        "char_start": start,
                        "char_end": end,
                    }
                )
                cursor = end + 2
        else:
            paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
            cursor = 0
            for p in paras:
                start = cursor
                end = start + len(p)
                sections.append(
                    {
                        "text": p,
                        "page_number": None,
                        "slide_index": None,
                        "section_heading": None,
                        "char_start": start,
                        "char_end": end,
                    }
                )
                cursor = end + 2
    else:
        paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
        cursor = 0
        for p in paras:
            start = cursor
            end = start + len(p)
            sections.append(
                {
                    "text": p,
                    "page_number": None,
                    "slide_index": None,
                    "section_heading": None,
                    "char_start": start,
                    "char_end": end,
                }
            )
            cursor = end + 2

    full_text = "\n\n".join([s["text"] for s in sections])
    if not sections:
        full_text = text
        sections = [
            {
                "text": full_text,
                "page_number": None,
                "slide_index": None,
                "section_heading": None,
                "char_start": 0,
                "char_end": len(full_text),
            }
        ]

    return {
        "text": full_text,
        "file_type": "md" if path.suffix.lower() == ".md" else "txt",
        "source_file": str(path),
        "filename": path.name,
        "document_title": None,
        "sections": sections,
    }


def load_file(path: str) -> Dict:
    """Load a file and return structured data including sections and metadata.

    Supported types: pdf, docx, pptx, csv, txt, md
    """
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        return read_pdf(p)
    if suffix in (".docx", ".doc"):
        return read_docx(p)
    if suffix in (".pptx", ".ppt"):
        return read_pptx(p)
    if suffix == ".csv":
        return read_csv(p)
    if suffix in (".md", ".txt"):
        return read_text_or_md(p)

    # fallback: return raw text
    raw = p.read_text(encoding="utf-8")
    return {
        "text": raw,
        "file_type": suffix.lstrip("."),
        "source_file": str(p),
        "filename": p.name,
        "document_title": None,
        "sections": [
            {
                "text": raw,
                "page_number": None,
                "slide_index": None,
                "section_heading": None,
                "char_start": 0,
                "char_end": len(raw),
            }
        ],
    }
